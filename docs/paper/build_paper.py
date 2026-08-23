from __future__ import annotations

import argparse
import csv
import hashlib
import json
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PAPER_DIR = Path(__file__).resolve().parent
FIGURE_DIR = PAPER_DIR / "figures"
APPENDIX_DIR = PAPER_DIR / "appendix"
TABLE_DIR = PAPER_DIR / "tables"
MARKDOWN_PATH = PAPER_DIR / "quant_strategy_paper_cn.md"
DOCX_PATH = PAPER_DIR / "quant_strategy_paper_cn.docx"
PDF_PATH = PAPER_DIR / "quant_strategy_paper_cn.pdf"
MANIFEST_PATH = PAPER_DIR / "reproducibility_manifest.json"
METRICS_PATH = (
    ROOT
    / "outputs"
    / "high_return_v2"
    / "next_open_rank_model_stock_focus_lev093_marketfilter_bench20260722_20220101_20260722"
    / "metrics.json"
)
EQUITY_PATH = METRICS_PATH.parent / "equity_curve.csv"
WEIGHTS_PATH = METRICS_PATH.parent / "rolling_feature_weights.csv"
RUN_CARD_PATH = ROOT / "outputs" / "high_return_v2" / "daily_run_card_20260722.json"
PRELIVE_PATH = ROOT / "outputs" / "high_return_v2" / "prelive_order_draft_20260722.json"
REGIME_PATH = ROOT / "outputs" / "high_return_v2" / "regime_shadow_tracking_summary.json"
BREADTH_PATH = (
    ROOT / "outputs" / "high_return_v2" / "dynamic_breadth_overlay_tracking_summary.json"
)
PROFILE_PATH = APPENDIX_DIR / "figure_data_profile.json"

TITLE = "动态风险约束与整手执行下的A股多因子趋势-回调组合策略研究"
SUBTITLE = "面向主板与创业板的滚动训练、次日开盘执行和实盘前置验证"
STATUS = "研究预印本 v0.1"
DATA_CUTOFF = "2026-07-22"
AUTHOR = "量化研究项目组"

FONT_CN = "Microsoft YaHei"
FONT_LATIN = "Calibri"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)
RISK_RED = RGBColor(155, 28, 28)
CAUTION = RGBColor(122, 90, 0)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def git_output(*args: str) -> str:
    result = subprocess.run(
        ["git", *args],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return result.stdout.strip()


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f}%"


def fmt_float(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def annual_returns() -> dict[str, float]:
    profile = read_json(PROFILE_PATH)
    return {str(key): float(value) for key, value in profile["annual_net_returns"].items()}


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_run_font(
    run,
    *,
    name: str = FONT_CN,
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    width.set(qn("w:type"), "dxa")
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(indent_dxa))
    indent.set(qn("w:type"), "dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(value))
        grid.append(grid_col)

    for row in table.rows:
        for cell, value in zip(row.cells, widths_dxa):
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(value))
            cell_width.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths_dxa: Sequence[int],
    *,
    numeric_columns: set[int] | None = None,
) -> None:
    data = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(value))
        set_run_font(run, size=8.5, bold=True)
    numeric_columns = numeric_columns or set()
    for row_values in data:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_run_font(run, size=8.5)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=MUTED)


def add_figure(doc: Document, filename: str, caption: str, alt_text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(FIGURE_DIR / filename), width=Inches(6.4))
    doc_properties = inline._inline.docPr
    doc_properties.set("title", caption)
    doc_properties.set("descr", alt_text)
    add_caption(doc, caption)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = False
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        tail = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(tail)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_bullet(doc: Document, text: str, *, numbered: bool = False) -> None:
    paragraph = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    set_run_font(run)


def add_callout(doc: Document, label: str, text: str, *, risk: bool = False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:cantSplit"))
    cell = table.cell(0, 0)
    shade_cell(cell, "FFF2F2" if risk else CALLOUT)
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    lead = paragraph.add_run(f"{label}：")
    set_run_font(lead, bold=True, color=RISK_RED if risk else DARK_BLUE)
    body = paragraph.add_run(text)
    set_run_font(body)
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_run_font(run, bold=True)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = FONT_CN
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    caption.font.size = Pt(9)
    caption.font.color.rgb = MUTED

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(0)
    run = header_paragraph.add_run("A股多因子组合策略研究 | 预印本 v0.1")
    set_run_font(run, size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = AUTHOR
    doc.core_properties.keywords = "A股; 多因子; 滚动训练; 次日开盘; 风险控制; 影子账户"
    doc.core_properties.comments = "研究预印本，非投资建议，不连接券商，不自动下单。"


def add_cover(doc: Document) -> None:
    for _ in range(3):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(14)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("QUANTITATIVE STRATEGY RESEARCH")
    set_run_font(run, name=FONT_LATIN, size=10, color=CAUTION, bold=True)
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run(TITLE)
    set_run_font(run, size=25, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run(SUBTITLE)
    set_run_font(run, size=13.5, color=MUTED)

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(8)
    run = metadata.add_run(f"{STATUS} | 数据截至 {DATA_CUTOFF}")
    set_run_font(run, size=11, bold=True, color=DARK_BLUE)

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(50)
    run = author.add_run(AUTHOR)
    set_run_font(run, size=10.5, color=MUTED)

    add_callout(
        doc,
        "研究声明",
        "本文记录可复现的历史研究和实盘前置验证状态。所有候选、目标仓位和影子账户均为研究用途，不构成投资建议，不连接券商，也不允许自动下单。",
        risk=True,
    )
    doc.add_page_break()


def add_abstract(doc: Document, metrics: dict) -> None:
    add_heading(doc, "摘要", 1)
    add_body(
        doc,
        "本文构建并审计一个面向A股主板与创业板的多因子股票组合研究系统。系统使用滚动横截面RankIC训练，在收盘后生成信号、下一交易日开盘调整，并将市场状态过滤、涨跌停与开盘跳空约束、交易成本、容量压力和一万元整手影子账户纳入同一证据链。数据面板截至2026年7月22日，共4,841,868条记录、4,600只证券和1,102个交易日期。",
    )
    add_body(
        doc,
        f"严格历史基线覆盖847个实现收益会话，总收益为{pct(metrics['total_return'])}，年化收益为{pct(metrics['annualized_return'])}，最大回撤为{pct(metrics['max_drawdown'])}，Sharpe-like为{fmt_float(metrics['sharpe_like'])}。但收益高度集中于2025年，其余三个日历分段均为负，说明历史正收益尚未形成跨年度稳定证据。多周期训练、纯趋势、纯均值回归、固定混合、单因子强势回调和日频即时减仓等挑战者均未通过预登记门槛。",
    )
    add_body(
        doc,
        "截至数据截止日，动态市场状态策略仅积累7/20个前瞻观察日，动态市场宽度覆盖为0/60，一万元整手影子账户尚未发生首笔真实前向模拟成交。因此本文将当前成果定义为研究预印本，而非最终实证论文或可直接投入实盘的策略说明书。",
    )
    keywords = doc.add_paragraph()
    lead = keywords.add_run("关键词：")
    set_run_font(lead, bold=True, color=DARK_BLUE)
    tail = keywords.add_run("A股；多因子选股；滚动训练；次日开盘；组合风控；回测过拟合；整手影子账户")
    set_run_font(tail)


def add_evidence_hierarchy(doc: Document, regime: dict, breadth: dict) -> None:
    add_heading(doc, "1 研究问题与证据边界", 1)
    add_body(
        doc,
        "动量研究表明中期赢家组合可能延续强势，而反转研究提示极端价格变化后可能出现反向修复。二者并不意味着任意趋势-回调公式都能产生可交易收益；在A股环境中，信号可见时间、涨跌停、停牌、开盘跳空、交易成本和100股整手约束会显著改变组合路径。本文不从单个高收益片段出发推断策略有效，而是提出四个可核验问题：历史基线是否为正、收益是否跨年度稳定、执行约束后是否仍可实现、前瞻样本是否足以支持实盘。",
    )
    add_heading(doc, "1.1 证据层级", 2)
    add_table(
        doc,
        ["证据层", "当前状态", "可支持的结论", "禁止外推"],
        [
            ("严格历史回测", "已完成", "历史路径、成本和回撤诊断", "未来收益承诺"),
            ("历史压力与失败实验", "已完成", "容量边界与假设证伪", "新样本alpha"),
            ("市场状态前瞻观察", f"{regime['valid_observation_count']}/{regime['target_days']}", "有限的增量运行证据", "成熟晋级"),
            ("动态宽度前瞻观察", f"{breadth['valid_observation_count']}/{breadth['target_valid_trade_days']}", "尚无有效观察日", "风险层上线"),
            ("一万元整手影子账户", "等待首个下一交易日执行", "执行流程已冻结", "真实成交绩效"),
            ("券商实盘", "未连接", "无", "自动下单"),
        ],
        [1550, 1500, 3030, 3280],
    )
    add_callout(
        doc,
        "核心边界",
        "历史回测、重叠前瞻统计和真实前向影子成交是三类不同证据。本文所有表格均保留其证据类别，旧版1000%以上收益曲线只作为偏差审计材料，不作为当前策略收益。",
        risk=True,
    )


def add_data_method(doc: Document) -> None:
    add_heading(doc, "2 数据、股票池与点时约束", 1)
    add_heading(doc, "2.1 数据面板", 2)
    add_body(
        doc,
        "研究面板覆盖2022年1月4日至2026年7月22日，包含date、symbol、open、high、low、close、volume、amount以及有限覆盖的主力净流入字段。每日流程优先读取 D:\\codex\\daily-market-data，在缺少最新数据时才检查备用抓取目录。2026年6月22日及之后的日数据来自统一标准化目录。",
    )
    add_table(
        doc,
        ["项目", "数值", "说明"],
        [
            ("面板记录数", "4,841,868", "长表记录"),
            ("证券数量", "4,600", "主板与创业板历史代码"),
            ("日期数量", "1,102", "面板内不同交易日期"),
            ("历史回测实现会话", "847", "2023-01-20至2026-07-22"),
            ("主力净量可用历史", "7个源会话", "仅作前向观察"),
        ],
        [1900, 1700, 5760],
        numeric_columns={1},
    )
    add_heading(doc, "2.2 股票池与排除规则", 2)
    add_bullet(doc, "正式研究股票池仅包含A股主板与创业板代码前缀；ETF、基金和其他板块不进入当前组合。")
    add_bullet(doc, "ST与*ST在点时5%涨跌停和交易状态未完整建模前，只能进入风险观察，不得占用模拟持仓。")
    add_bullet(doc, "需要人工复核的证券不进入正式观察目标；异常日收益绝对值超过22%的记录在训练/回测清洗中受到限制。")
    add_bullet(doc, "当前未使用1500亿元市值排除条件；大市值证券能否入选由模型排序、趋势状态和执行约束共同决定。")
    add_heading(doc, "2.3 标签可见性与未来函数防护", 2)
    add_body(
        doc,
        "一日实现收益标签定义为 r(i,t)=Open(i,t+2)/Open(i,t+1)-1。信号在t日收盘后可见，最早在t+1开盘成交，收益在t+2开盘完成。对持有期h，信号位置i可用于训练的最后一个成熟标签只能位于i-h-1；训练切片右端为i-h。当前生产冠军使用h=1、252个成熟训练日、每20个交易日重训。",
    )
    add_callout(
        doc,
        "方法限制",
        "当前实现收益口径是一交易会话的次日开盘到后续开盘，与用户偏好的数周至数月持有期不完全一致。延长训练标签到1/3/5日但不改变实际持有期的挑战者已被拒绝；后续需要重新设计持有期一致的独立策略袖套。",
    )


def add_features_training(doc: Document) -> None:
    add_heading(doc, "3 因子体系与滚动训练", 1)
    add_heading(doc, "3.1 冻结的生产特征", 2)
    add_table(
        doc,
        ["特征族", "代表特征", "经济含义", "主要风险"],
        [
            ("趋势", "5/20/60日动量、20日突破、距MA20", "价格延续与趋势位置", "追高和状态切换"),
            ("反转", "5日反转、日内收益、收盘位置", "短期过度反应修复", "单边下跌暴露"),
            ("风险与流动性", "20日波动、20日成交额", "波动与可交易性", "流动性倾斜和容量"),
            ("趋势-回调交互", "20/60日强势回调、突破后回调", "中期强势中的短期回撤", "变量重叠与重复暴露"),
            ("反追高", "日内反追高、流动性回调", "避免短线高潮买入", "可能错过强趋势"),
        ],
        [1450, 2750, 2480, 2680],
    )
    add_body(
        doc,
        "每个交易日先对特征做横截面百分位排名，再计算特征与未来次日开盘收益的Spearman RankIC。训练窗口内平均IC按绝对值归一化形成因子权重。生产模型使用显式的15特征白名单，避免研究特征因代码新增而自动进入正式训练。权重只解释打分倾向，不能等同于逐因子利润归因。",
    )
    add_heading(doc, "3.2 组合构建", 2)
    add_bullet(doc, "训练窗口252个成熟交易日；每20个交易日重训一次。")
    add_bullet(doc, "每5个交易日调仓；选取综合得分前40只股票。")
    add_bullet(doc, "基础总仓位上限0.93；单票上限2.9%；仅做多。")
    add_bullet(doc, "组合目标权重先由模型决定，再乘正式市场风险暴露；实验性宽度过滤不改变生产权重。")
    add_figure(
        doc,
        "figure_3_rolling_factor_weights.png",
        "图3 滚动因子权重热力图",
        "从2023年至2026年各重训日期的15个因子归一化RankIC权重。红色为正，蓝色为负。",
    )
    add_body(
        doc,
        "图3显示因子权重随时间明显变化，且短期动量与5日反转存在严格镜像关系。该结构会重复表达同一底层价格变化，是当前模型解释性和稳健性的主要限制之一。纯趋势、纯反转和固定混合拆分均未复现冠军收益，说明冠军来自多特征联合排序，而不是某个单一可稳定外推的因子。",
    )


def add_execution_risk(doc: Document) -> None:
    add_heading(doc, "4 成交、成本与组合风控", 1)
    add_heading(doc, "4.1 次日开盘成交合同", 2)
    add_table(
        doc,
        ["项目", "生产回测设定", "实盘前置影子设定"],
        [
            ("信号时间", "收盘后", "收盘后生成草稿"),
            ("最早成交", "下一交易日开盘", "下一实际交易日开盘"),
            ("组合成本", "佣金1.0bp + 冲击0.7bp", "佣金3bp且最低5元；滑点5bp"),
            ("卖出税费", "组合成本统一近似", "卖出印花税5bp"),
            ("整手", "权重模拟", "100股"),
            ("开盘跳空", "买入跳空超过6%受限", "买入跳空超过6%阻塞"),
            ("涨跌停", "按板块默认限制和0.995缓冲", "同口径并记录阻塞"),
        ],
        [1800, 3480, 4080],
    )
    add_heading(doc, "4.2 正式市场风险层", 2)
    add_body(
        doc,
        "正式风险层以510300收盘价为市场代理。基准高于120日均线且20日跌幅高于-8%时，风险暴露为100%；基准低于MA120时为60%；20日跌幅小于等于-8%时为0%。该暴露只在已知收盘数据上计算，并乘到基础策略目标仓位。生产冠军仍采用每5日定期调仓；历史测试显示在非调仓日立即减仓会损失反弹收益，因此日频只减不增规则被拒绝。",
    )
    add_figure(
        doc,
        "figure_2_calendar_exposure.png",
        "图2 年度收益与市场风险目标暴露",
        "左图为各日历年度复合净收益，右图为0%、60%和100%市场风险目标暴露出现的交易会话数。",
    )
    add_callout(
        doc,
        "重要区分",
        "市场风险暴露是目标上限，不是实际已部署资金。实际仓位必须读取gross_exposure；因开盘限制、非调仓日和整手约束，二者可能不同。",
    )


def add_results(doc: Document, metrics: dict, annual: dict[str, float]) -> None:
    add_heading(doc, "5 严格历史基线结果", 1)
    add_table(
        doc,
        ["指标", "结果", "解释"],
        [
            ("总收益", pct(metrics["total_return"]), "初始资金至期末净值"),
            ("年化收益", pct(metrics["annualized_return"]), "按完整净收益序列计算"),
            ("最大回撤", pct(metrics["max_drawdown"]), "历史峰值至谷值"),
            ("Sharpe-like", fmt_float(metrics["sharpe_like"]), "未声明为统计显著性检验"),
            ("平均换手", pct(metrics["avg_turnover"]), "每日权重绝对变动之和"),
            ("平均实际总仓位", pct(metrics["avg_gross_exposure"]), "约40只持仓"),
            ("平均市场风险目标", pct(metrics["avg_market_exposure_target"]), "不等于实际仓位"),
            ("回测会话", str(metrics["trade_days"]), "实现净收益的交易会话"),
        ],
        [2350, 1850, 5160],
        numeric_columns={1},
    )
    add_figure(
        doc,
        "figure_1_equity_drawdown.png",
        "图1 严格历史基线净值与回撤",
        "2023年1月20日至2026年7月22日的归一化净值和从历史高点计算的回撤。",
    )
    add_heading(doc, "5.1 年度稳定性", 2)
    add_table(
        doc,
        ["日历分段", "净收益", "判断"],
        [
            ("2023", pct(annual["2023"]), "负收益"),
            ("2024", pct(annual["2024"]), "负收益"),
            ("2025", pct(annual["2025"]), "唯一正收益年份"),
            ("2026截至7月22日", pct(annual["2026"]), "负收益"),
        ],
        [3100, 1900, 4360],
        numeric_columns={1},
    )
    add_body(
        doc,
        "全周期收益为正，但四个日历分段中只有2025年为正，且该年度贡献了全部正的日历对数收益。风险开启状态贡献了82.84%的全周期净对数收益，但该状态在2023年和2026年仍为负。因此，7.55%的历史年化收益不能被外推为稳定年度收益，更不能支持500%或更高的实盘收益目标。",
    )
    add_heading(doc, "5.2 执行与风险审计", 2)
    add_bullet(doc, "历史回测记录6个出现开盘约束的会话，累计受限证券46个；没有启用组合容量限制的正式生产基线。")
    add_bullet(doc, "有2个市场目标为0但处于非调仓日的会话保留了仓位；它们属于定期调仓政策的风险响应延迟，不是已证实的涨跌停卖出阻塞。")
    add_bullet(doc, "liquidity_20在全部重训点均为负权重，表明模型持续偏向较低流动性标的；扩大资金或集中度前必须先做容量压力。")


def add_failures(doc: Document) -> None:
    add_heading(doc, "6 失败实验与反向证据", 1)
    add_body(
        doc,
        "失败实验是当前研究可信度的重要组成部分。所有下表候选均使用已观察历史，其作用是证伪假设而不是提供新的样本外收益。未通过一级门槛后不继续调参救回，避免在同一区间重复寻找看似成功的解释。",
    )
    add_table(
        doc,
        ["候选", "总收益", "最大回撤", "决定", "主要原因"],
        [
            ("1/3/5日多周期训练", "21.24%", "-28.34%", "拒绝", "收益与Sharpe均低于冠军"),
            ("纯趋势袖套", "-28.25%", "-28.83%", "拒绝", "一日标签下趋势方向失效"),
            ("纯均值回归袖套", "8.75%", "-30.82%", "拒绝", "收益低且回撤更差"),
            ("固定50/50混合", "-9.03%", "-22.49%", "拒绝", "低相关不能挽救负收益袖套"),
            ("强势回调单因子", "-51.05%", "-54.13%", "拒绝", "正RankIC未转化为可成交组合收益"),
            ("趋势过滤回调", "20.63%", "-20.04%", "拒绝", "回撤改善但收益和Sharpe未超冠军"),
            ("每日只减仓风险响应", "17.33%", "-29.79%", "拒绝", "等待恢复期间机会损失显著"),
        ],
        [2140, 1200, 1250, 950, 3820],
        numeric_columns={1, 2},
    )
    add_heading(doc, "6.1 RankIC与组合收益分离", 2)
    add_body(
        doc,
        "强势回调因子的平均一日RankIC为0.0166，正IC比例为57.96%，43个重训点均得到正权重，但执行后的组合在每个日历年度都亏损。该结果说明，全股票池的弱正排序关系不等于前40名组合在次日开盘、市场过滤和成本约束下可盈利。后续因子必须同时通过相关性诊断和可成交组合门槛。",
    )
    add_heading(doc, "6.2 旧高收益曲线的处理", 2)
    add_callout(
        doc,
        "失败审计规则",
        "项目历史上出现过1000%至10000%以上的曲线。凡未完成点时股票池、未来函数、复权、成交约束、幸存者偏差和独立样本审计的结果，统一标记为legacy bias evidence，不得进入当前策略收益表。",
        risk=True,
    )


def add_capacity(doc: Document) -> None:
    add_heading(doc, "7 容量压力与小账户整手执行", 1)
    add_heading(doc, "7.1 历史容量压力", 2)
    add_body(
        doc,
        "容量压力按信号日已知的20日成交额中位数计算，单日买卖最多使用该金额的5%。无容量控制路径先精确复现冠军，然后再比较不同资金规模。该测试只评估执行损伤，不提供新增alpha。",
    )
    add_table(
        doc,
        ["资金规模", "总收益", "最大回撤", "仓位保留", "成交填充", "门槛"],
        [
            ("1000万元", "29.66%", "-28.56%", "100.68%", "99.57%", "通过"),
            ("5000万元", "30.19%", "-26.17%", "96.10%", "95.56%", "通过"),
            ("1亿元", "26.24%", "-22.11%", "81.13%", "84.50%", "临界通过"),
            ("5亿元", "4.01%", "-8.13%", "29.04%", "40.35%", "失败"),
        ],
        [1600, 1350, 1450, 1600, 1550, 1810],
        numeric_columns={1, 2, 3, 4},
    )
    add_body(
        doc,
        "1亿元是预登记场景中的最大历史通过值，但其仓位保留仅略高于80%底线，因此只能作为粗略历史上界。1000万和5000万场景略高于无容量控制路径，是部分成交改变持仓路径产生的噪声，不能解释为容量限制创造了收益。",
    )
    add_heading(doc, "7.2 一万元实盘前置影子账户", 2)
    add_body(
        doc,
        "实盘前置账户使用10,000元初始资金、100股整手、先卖后买、最低5元佣金、卖出税费和滑点。整手分配使用输入顺序无关的背包式算法，避免候选表先后顺序静默占用现金。影子账户只能在下一实际交易日执行前一交易日草稿；首日没有旧草稿时必须等待，禁止用当日开盘价回填。",
    )
    add_table(
        doc,
        ["截至日期", "基础上限", "市场暴露", "有效上限", "目标市值", "执行样本"],
        [("2026-07-22", "93.0%", "60.0%", "55.8%", "5,253元", "0个完整前向成交日")],
        [1550, 1400, 1400, 1400, 1550, 2060],
        numeric_columns={1, 2, 3, 4},
    )
    add_callout(
        doc,
        "当前状态",
        "2026年7月22日草稿虽然给出5个整手目标，但仍是manual_review_only。trade_instruction=false、broker_order_allowed=false，广发证券连接和自动下单均关闭。",
        risk=True,
    )


def add_forward_validation(doc: Document, regime: dict, breadth: dict) -> None:
    add_heading(doc, "8 前瞻验证与晋级规则", 1)
    add_table(
        doc,
        ["观察项目", "登记起点", "当前进度", "成熟门槛", "生产影响"],
        [
            ("市场状态动态策略", "2026-07-14附近", f"{regime['valid_observation_count']}/{regime['target_days']}", "20个有效交易日", "实验提示"),
            ("动态市场宽度覆盖", breadth["observation_start_date"], f"{breadth['valid_observation_count']}/{breadth['target_valid_trade_days']}", "60个有效交易日", "无"),
            ("机构吸筹观察", "2026-07-20", "0/80个5日完成样本", "80个主周期样本", "无"),
            ("一万元整手影子账户", "首个旧草稿之后", "等待首个执行日", "连续前向成交与对账", "人工复核"),
        ],
        [1900, 1750, 1900, 2020, 1790],
    )
    add_bullet(doc, "参数进入未见样本后冻结；继续调参必须使用新登记编号和新验证起点。")
    add_bullet(doc, "重复运行同一日期必须幂等，不能增加虚假的独立观察日。")
    add_bullet(doc, "历史Pareto占优只用于发现挑战者，不能作为晋级证据；任何晋级都要求人工复核。")
    add_bullet(doc, "自动晋级、跨策略族晋级和券商自动下单保持关闭。")


def add_discussion(doc: Document) -> None:
    add_heading(doc, "9 讨论", 1)
    add_heading(doc, "9.1 可以得到的结论", 2)
    add_bullet(doc, "严格时间标签修复后，生产冠军在已观察历史中保持正收益，说明结果不是由最近一条未成熟标签直接造成。")
    add_bullet(doc, "市场风险过滤、开盘约束和成本已进入回测，容量与整手执行已有独立审计路径。")
    add_bullet(doc, "预登记、顺序停止、失败记录和前瞻观察账本已经形成抑制过拟合的工程框架。")
    add_heading(doc, "9.2 不能得到的结论", 2)
    add_bullet(doc, "不能把27.71%的历史总收益解释为稳定年化，更不能推出至少500%的未来收益。")
    add_bullet(doc, "不能将2025年的强势区间视为已证明可复制；其他三个日历分段均为负。")
    add_bullet(doc, "不能将正RankIC、候选名单或模型目标权重等同于真实可成交利润。")
    add_bullet(doc, "不能把尚未成熟的市场状态、宽度、资金流和形态观察直接用于正式仓位。")
    add_heading(doc, "9.3 与经典研究的关系", 2)
    add_body(
        doc,
        "本研究同时借鉴中期动量、长期过度反应与多因子风险解释，但并不假定这些现象会在A股一日持有期中直接复制。White的数据窥探检验、Harvey等人的多重检验警告以及Bailey等人的回测过拟合概率框架共同支持本项目的预登记和前瞻观察制度。当前项目尚未计算White Reality Check、Deflated Sharpe Ratio或组合对称交叉验证概率，因此仍缺少正式的多重试验校正。",
    )


def add_limitations_roadmap(doc: Document) -> None:
    add_heading(doc, "10 局限与后续研究", 1)
    add_heading(doc, "10.1 主要局限", 2)
    limitations = [
        "收益稳定性不足：四个日历分段只有2025年为正，未建立跨年度和跨状态稳定性。",
        "持有期不一致：实际损益仍是一交易会话，尚未验证数周至数月目标。",
        "因子解释不充分：动量5与反转5严格镜像，交互特征存在底层变量重叠。",
        "风险模型较简化：当前没有行业/市值中性化，也未使用协方差矩阵进行组合波动估计。",
        "基准比较不足：尚未形成统一的相对沪深300超额收益、beta和显著性表。",
        "数据边界：资金流仅有短期前向历史，ST点时涨跌停和异常交易状态仍不完整。",
        "前瞻样本不足：市场状态、宽度和整手账户均未达到成熟门槛。",
        "交易落地差异：回测权重模型与一万元整手账户在最低佣金、持仓离散化和现金利用率上存在显著差异。",
    ]
    for item in limitations:
        add_bullet(doc, item)
    add_heading(doc, "10.2 优先级路线图", 2)
    roadmap = [
        "完成一万元整手影子账户的首个真实前向执行，并连续积累至少60个交易日，逐日对账目标、成交、费用、现金和净值。",
        "按冻结登记完成市场状态20日和动态宽度60日观察，不在观察期内修改门槛。",
        "建立持有期一致的3/5/10/20日策略袖套，训练标签、持仓更新和退出价格使用同一周期。",
        "对因子做行业与市值中性化对照，并消除完全镜像或高度同源的重复特征。",
        "补充沪深300相对绩效、beta、信息比率、年度与状态分层，以及White Reality Check、PBO/CSCV和Deflated Sharpe Ratio。",
        "在前瞻证据成熟后再设计小额人工实盘门槛；券商接口、密钥管理、撤单与风控熔断另行审核。",
    ]
    for item in roadmap:
        add_bullet(doc, item, numbered=True)


def add_conclusion(doc: Document) -> None:
    add_heading(doc, "11 结论", 1)
    add_body(
        doc,
        "本项目已经从单次高收益回测推进到具备严格标签清除、冻结特征白名单、次日开盘约束、市场风险层、容量压力、失败实验账本和整手影子账户的研究系统。严格历史冠军取得27.71%的总收益和-28.26%的最大回撤，但收益集中于单一年份，挑战者未形成更优且稳定的替代方案，前瞻样本也尚未成熟。",
    )
    add_body(
        doc,
        "因此当前最合理的决策是保留冠军作为研究基准，继续冻结并积累前向证据，而不是扩大收益承诺或立即接入自动实盘。本文作为活文档，可在每次月度研究节点更新数据截止日、前瞻样本、影子账户净值和失败实验，同时保持历史版本与复现清单。",
    )


def add_references(doc: Document) -> None:
    add_heading(doc, "参考文献", 1)
    references = [
        (
            "Jegadeesh, N., & Titman, S. (1993). Returns to Buying Winners and Selling Losers: Implications for Stock Market Efficiency. Journal of Finance, 48(1), 65-91.",
            "https://doi.org/10.1111/j.1540-6261.1993.tb04702.x",
        ),
        (
            "De Bondt, W. F. M., & Thaler, R. (1985). Does the Stock Market Overreact? Journal of Finance, 40(3), 793-805.",
            "https://doi.org/10.1111/j.1540-6261.1985.tb05004.x",
        ),
        (
            "Fama, E. F., & French, K. R. (1993). Common Risk Factors in the Returns on Stocks and Bonds. Journal of Financial Economics, 33(1), 3-56.",
            "https://doi.org/10.1016/0304-405X(93)90023-5",
        ),
        (
            "White, H. (2000). A Reality Check for Data Snooping. Econometrica, 68(5), 1097-1126.",
            "https://doi.org/10.1111/1468-0262.00152",
        ),
        (
            "Harvey, C. R., Liu, Y., & Zhu, H. (2016). ... and the Cross-Section of Expected Returns. Review of Financial Studies, 29(1), 5-68.",
            "https://doi.org/10.1093/rfs/hhv059",
        ),
        (
            "Bailey, D. H., Borwein, J. M., Lopez de Prado, M., & Zhu, Q. J. (2015). The Probability of Backtest Overfitting.",
            "https://doi.org/10.2139/ssrn.2326253",
        ),
    ]
    for citation, url in references:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(citation + " ")
        set_run_font(run, size=9.5)
        add_hyperlink(paragraph, url, url)


def add_reproducibility(doc: Document, run_card: dict) -> None:
    doc.add_page_break()
    add_heading(doc, "附录A 复现说明", 1)
    add_body(
        doc,
        "论文源文件、图表生成脚本、Word构建脚本和机器可读清单位于docs/paper。图表只读取正式生产冠军的equity_curve.csv与rolling_feature_weights.csv；正文关键数字读取metrics.json、每日运行卡和冻结观察摘要。",
    )
    add_body(
        doc,
        "复现命令：先运行python docs/paper/build_figures.py，再使用Codex工作区Python运行python docs/paper/build_paper.py。Word与PDF需通过可用的Microsoft Word或LibreOffice渲染并逐页检查。完整输入路径与SHA256记录在reproducibility_manifest.json。",
    )
    add_table(
        doc,
        ["复现项目", "记录"],
        [
            ("数据截止日", DATA_CUTOFF),
            ("每日流程状态", run_card["run_status"]),
            ("每日流程步骤", f"{run_card['command_count']}个成功步骤"),
            ("运行卡测试", run_card["verification"]["tests"]),
            ("Git提交", git_output("rev-parse", "HEAD")[:12]),
            ("工作区状态", "包含未提交研究改动；详见manifest"),
            ("自动交易", "关闭"),
        ],
        [2500, 6860],
    )


def build_docx(
    metrics: dict,
    run_card: dict,
    prelive: dict,
    regime: dict,
    breadth: dict,
) -> None:
    del prelive
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_abstract(doc, metrics)
    add_evidence_hierarchy(doc, regime, breadth)
    add_data_method(doc)
    add_features_training(doc)
    add_execution_risk(doc)
    add_results(doc, metrics, annual_returns())
    add_failures(doc)
    add_capacity(doc)
    add_forward_validation(doc, regime, breadth)
    add_discussion(doc)
    add_limitations_roadmap(doc)
    add_conclusion(doc)
    add_references(doc)
    add_reproducibility(doc, run_card)
    doc.save(DOCX_PATH)


def markdown_table(headers: Sequence[str], rows: Iterable[Sequence[str]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend("| " + " | ".join(str(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def build_markdown(metrics: dict, run_card: dict, regime: dict, breadth: dict) -> None:
    annual = annual_returns()
    evidence = markdown_table(
        ["证据层", "当前状态", "可支持的结论", "禁止外推"],
        [
            ("严格历史回测", "已完成", "历史路径、成本和回撤诊断", "未来收益承诺"),
            ("市场状态前瞻观察", f"{regime['valid_observation_count']}/{regime['target_days']}", "有限增量证据", "成熟晋级"),
            ("动态宽度前瞻观察", f"{breadth['valid_observation_count']}/{breadth['target_valid_trade_days']}", "尚无有效观察日", "风险层上线"),
            ("一万元整手影子账户", "等待首个执行日", "流程已冻结", "真实成交绩效"),
            ("券商实盘", "未连接", "无", "自动下单"),
        ],
    )
    results = markdown_table(
        ["指标", "结果"],
        [
            ("总收益", pct(metrics["total_return"])),
            ("年化收益", pct(metrics["annualized_return"])),
            ("最大回撤", pct(metrics["max_drawdown"])),
            ("Sharpe-like", fmt_float(metrics["sharpe_like"])),
            ("平均换手", pct(metrics["avg_turnover"])),
            ("平均实际总仓位", pct(metrics["avg_gross_exposure"])),
            ("会话数", str(metrics["trade_days"])),
        ],
    )
    yearly = markdown_table(
        ["分段", "净收益"],
        [(year, pct(value)) for year, value in annual.items()],
    )
    failures = markdown_table(
        ["候选", "总收益", "最大回撤", "决定"],
        [
            ("1/3/5日多周期训练", "21.24%", "-28.34%", "拒绝"),
            ("纯趋势袖套", "-28.25%", "-28.83%", "拒绝"),
            ("纯均值回归袖套", "8.75%", "-30.82%", "拒绝"),
            ("固定50/50混合", "-9.03%", "-22.49%", "拒绝"),
            ("强势回调单因子", "-51.05%", "-54.13%", "拒绝"),
            ("趋势过滤回调", "20.63%", "-20.04%", "拒绝"),
            ("每日只减仓风险响应", "17.33%", "-29.79%", "拒绝"),
        ],
    )
    text = f"""# {TITLE}

{SUBTITLE}

**{STATUS} | 数据截至 {DATA_CUTOFF}**  
**作者：{AUTHOR}**

> **研究声明：** 本文记录可复现的历史研究和实盘前置验证状态，不构成投资建议；券商连接和自动下单均关闭。

## 摘要

本文构建并审计一个面向A股主板与创业板的多因子股票组合研究系统。系统使用滚动横截面RankIC训练，在收盘后生成信号、下一交易日开盘调整，并将市场状态过滤、涨跌停与开盘跳空约束、交易成本、容量压力和一万元整手影子账户纳入同一证据链。数据面板截至2026年7月22日，共4,841,868条记录、4,600只证券和1,102个交易日期。

严格历史基线覆盖847个实现收益会话，总收益为{pct(metrics['total_return'])}，年化收益为{pct(metrics['annualized_return'])}，最大回撤为{pct(metrics['max_drawdown'])}，Sharpe-like为{fmt_float(metrics['sharpe_like'])}。但收益高度集中于2025年，其余三个日历分段均为负。截至数据截止日，市场状态策略仅积累{regime['valid_observation_count']}/{regime['target_days']}个前瞻观察日，动态宽度覆盖为{breadth['valid_observation_count']}/{breadth['target_valid_trade_days']}，一万元整手影子账户尚未发生首笔真实前向模拟成交。因此本文属于研究预印本，不是最终实证结论。

**关键词：** A股；多因子选股；滚动训练；次日开盘；组合风控；回测过拟合；整手影子账户

## 1. 研究问题与证据边界

动量与反转是经典收益预测现象，但不能脱离A股点时数据、交易限制和成本直接照搬。本文检验历史基线、跨年度稳定性、执行可实现性和前瞻证据成熟度。

{evidence}

旧版1000%以上收益曲线统一归入偏差与失败审计，不作为当前策略收益。

## 2. 数据与方法

- 面板区间：2022-01-04至2026-07-22；4,841,868条记录；4,600只证券；1,102个日期。
- 股票池：A股主板与创业板；不做ETF/基金；ST与*ST不进入模拟持仓。
- 标签：`Open(t+2) / Open(t+1) - 1`；信号在`t`日收盘后生成，最早在`t+1`开盘成交。
- 训练：252个成熟交易日；每20日重训；每5日调仓；前40只；单票上限2.9%；基础总仓位上限0.93。
- 成本：佣金1.0bp + 冲击0.7bp；开盘跳空买入阈值6%；涨跌停缓冲0.995。
- 风险：510300高于MA120且20日跌幅高于-8%时100%，低于MA120时60%，20日跌幅不高于-8%时0%。

![图1 严格历史基线净值与回撤](figures/figure_1_equity_drawdown.png)

## 3. 因子与滚动训练

生产白名单包含15个价格、波动、流动性和趋势-回调交互特征。横截面百分位特征与未来次日开盘收益计算Spearman RankIC，252日平均IC按绝对值归一化形成权重。权重仅表示打分倾向，不是因子利润归因。

![图3 滚动因子权重热力图](figures/figure_3_rolling_factor_weights.png)

短期动量与5日反转严格镜像，交互项也存在底层变量重叠。纯趋势、纯均值回归和固定混合均未复现冠军收益。

## 4. 历史结果

{results}

![图2 年度收益与市场风险目标暴露](figures/figure_2_calendar_exposure.png)

{yearly}

全周期为正，但只有2025年为正。历史年化不能外推为稳定收益，更不能支持至少500%的未来收益目标。

## 5. 失败实验

{failures}

强势回调因子平均RankIC为0.0166且正IC比例57.96%，但执行组合在每个日历年度均亏损。这说明弱正排序关系不等于可成交组合利润。

## 6. 容量与整手执行

容量压力显示1亿元是预登记历史场景中的最大临界通过值，5亿元失败。一万元实盘前置账户使用100股整手、最低5元佣金、卖出税费与滑点；截至2026-07-22，有效仓位上限55.8%、目标市值5,253元，但仍为人工复核草稿，尚无完整前向成交日。

## 7. 前瞻验证

- 市场状态策略：{regime['valid_observation_count']}/{regime['target_days']}。
- 动态宽度覆盖：{breadth['valid_observation_count']}/{breadth['target_valid_trade_days']}。
- 机构吸筹观察：0/80个5日完成样本。
- 整手影子账户：等待首个下一交易日执行。
- 自动晋级和券商自动下单：关闭。

## 8. 局限与路线图

主要局限包括跨年度稳定性不足、持有期不一致、同源因子重复、缺少行业/市值中性化、缺少协方差风险模型、未完成基准超额收益与多重试验校正、资金流历史过短以及前瞻样本未成熟。

下一步优先完成整手影子账户和冻结观察门槛，再设计持有期一致的3/5/10/20日策略袖套，补充行业/市值中性化、沪深300相对绩效、White Reality Check、PBO/CSCV和Deflated Sharpe Ratio。

## 9. 结论

严格历史冠军取得27.71%的总收益和-28.26%的最大回撤，但收益集中于单一年份，挑战者未形成更优且稳定的替代方案，前瞻样本也尚未成熟。当前应保留冠军作为研究基准并继续积累独立证据，不应扩大收益承诺或立即接入自动实盘。

## 参考文献

1. [Jegadeesh & Titman (1993), Returns to Buying Winners and Selling Losers](https://doi.org/10.1111/j.1540-6261.1993.tb04702.x)
2. [De Bondt & Thaler (1985), Does the Stock Market Overreact?](https://doi.org/10.1111/j.1540-6261.1985.tb05004.x)
3. [Fama & French (1993), Common Risk Factors in the Returns on Stocks and Bonds](https://doi.org/10.1016/0304-405X(93)90023-5)
4. [White (2000), A Reality Check for Data Snooping](https://doi.org/10.1111/1468-0262.00152)
5. [Harvey, Liu & Zhu (2016), ... and the Cross-Section of Expected Returns](https://doi.org/10.1093/rfs/hhv059)
6. [Bailey et al. (2015), The Probability of Backtest Overfitting](https://doi.org/10.2139/ssrn.2326253)

## 复现

- 每日运行卡：`run_status={run_card['run_status']}`，{run_card['verification']['tests']}。
- Git提交：`{git_output('rev-parse', 'HEAD')}`。
- 构建：`python docs/paper/build_figures.py`，然后使用Codex工作区Python运行`python docs/paper/build_paper.py`。
- 完整输入路径和SHA256见`reproducibility_manifest.json`。
"""
    MARKDOWN_PATH.write_text(text, encoding="utf-8")


def build_manifest() -> dict:
    tracked_inputs = [
        ROOT / "train_next_open_rank_model.py",
        ROOT / "market_risk.py",
        ROOT / "execution_rules.py",
        ROOT / "prelive_order_draft.py",
        ROOT / "prelive_account_shadow.py",
        ROOT / "configs" / "prelive_gf_manual.yaml",
        METRICS_PATH,
        EQUITY_PATH,
        WEIGHTS_PATH,
        RUN_CARD_PATH,
        PRELIVE_PATH,
        REGIME_PATH,
        BREADTH_PATH,
        PROFILE_PATH,
        PAPER_DIR / "build_figures.py",
        PAPER_DIR / "build_paper.py",
    ]
    files = {}
    for path in tracked_inputs:
        files[str(path.relative_to(ROOT))] = {
            "path": str(path.resolve()),
            "size_bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
    data_panel = ROOT / "data_panel_history_main_chinext_20220101_20260722.parquet"
    manifest = {
        "schema_version": 1,
        "paper_status": "research_preprint_v0.1",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "data_cutoff": DATA_CUTOFF,
        "repository_root": str(ROOT),
        "git_commit": git_output("rev-parse", "HEAD"),
        "git_branch": git_output("branch", "--show-current"),
        "git_status_porcelain": git_output("status", "--short"),
        "data_panel": {
            "path": str(data_panel.resolve()),
            "size_bytes": data_panel.stat().st_size,
            "row_count": 4_841_868,
            "symbol_count": 4_600,
            "date_count": 1_102,
            "date_min": "2022-01-04",
            "date_max": DATA_CUTOFF,
            "note": "Large panel hash intentionally omitted; authoritative daily source and run-card hashes are retained separately.",
        },
        "inputs": files,
        "deliverables": {},
        "evidence_boundary": {
            "historical_backtest": True,
            "fresh_out_of_sample_alpha_complete": False,
            "exact_account_forward_execution_days": 0,
            "trade_instruction": False,
            "broker_order_allowed": False,
        },
    }
    return manifest


def finalize_manifest() -> None:
    manifest = read_json(MANIFEST_PATH) if MANIFEST_PATH.exists() else build_manifest()
    deliverables = {}
    for path in [MARKDOWN_PATH, DOCX_PATH, PDF_PATH]:
        if path.exists():
            deliverables[path.name] = {
                "path": str(path.resolve()),
                "size_bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
    manifest["deliverables"] = deliverables
    manifest["manifest_finalized_at"] = datetime.now().isoformat(timespec="seconds")
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the Chinese strategy research preprint.")
    parser.add_argument("--finalize-manifest", action="store_true")
    args = parser.parse_args()
    if args.finalize_manifest:
        finalize_manifest()
        return

    metrics = read_json(METRICS_PATH)
    run_card = read_json(RUN_CARD_PATH)
    prelive = read_json(PRELIVE_PATH)
    regime = read_json(REGIME_PATH)
    breadth = read_json(BREADTH_PATH)
    required_figures = [
        FIGURE_DIR / "figure_1_equity_drawdown.png",
        FIGURE_DIR / "figure_2_calendar_exposure.png",
        FIGURE_DIR / "figure_3_rolling_factor_weights.png",
    ]
    missing = [str(path) for path in required_figures if not path.exists()]
    if missing:
        raise FileNotFoundError(f"build figures before the paper: {missing}")
    PAPER_DIR.mkdir(parents=True, exist_ok=True)
    APPENDIX_DIR.mkdir(parents=True, exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    build_markdown(metrics, run_card, regime, breadth)
    build_docx(metrics, run_card, prelive, regime, breadth)
    MANIFEST_PATH.write_text(
        json.dumps(build_manifest(), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    finalize_manifest()
    print(str(DOCX_PATH))


if __name__ == "__main__":
    main()
